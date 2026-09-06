import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_placeholder_box(doc, placeholder_text, caption_text=""):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border teal
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:left w:val="single" w:sz="36" w:space="0" w:color="0D9488"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(f"🖼️  {placeholder_text}")
    run.bold = True
    run.font.name = 'Calibri'
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(13, 148, 136) # Teal
    
    if caption_text:
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(4)
        run2 = p2.add_run(caption_text)
        run2.italic = True
        run2.font.name = 'Calibri'
        run2.font.size = Pt(9.5)
        run2.font.color.rgb = RGBColor(71, 85, 105) # Slate
    
    # Add empty paragraph after table for spacing
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def build_word_document(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    doc = Document()
    
    # Page Margins 1 inch
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    lines = content.split('\n')
    i = 0
    n = len(lines)
    
    in_code_block = False
    code_lines = []
    
    while i < n:
        line = lines[i].rstrip()
        
        # Handle Fenced Code Blocks
        if line.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lines = []
            else:
                in_code_block = False
                # Write code block as styled callout box
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_background(cell, "0F172A") # Slate 900
                set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.1
                code_str = "\n".join(code_lines)
                run = p.add_run(code_str)
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(56, 189, 248) # Sky blue
                
                p_after = doc.add_paragraph()
                p_after.paragraph_format.space_after = Pt(6)
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Handle Empty Lines
        if not line.strip():
            i += 1
            continue
            
        # Handle Headings
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[2:].strip())
            run.bold = True
            run.font.name = 'Segoe UI'
            run.font.size = Pt(22)
            run.font.color.rgb = RGBColor(15, 23, 42) # Dark Slate
            i += 1
            continue
            
        if line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[3:].strip())
            run.bold = True
            run.font.name = 'Segoe UI'
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(13, 148, 136) # Teal accent
            i += 1
            continue
            
        if line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[4:].strip())
            run.bold = True
            run.font.name = 'Segoe UI'
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(30, 41, 59) # Slate
            i += 1
            continue
            
        if line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[5:].strip())
            run.bold = True
            run.font.name = 'Segoe UI'
            run.font.size = Pt(11.5)
            run.font.color.rgb = RGBColor(71, 85, 105)
            i += 1
            continue

        # Handle Placeholders & Captions
        if line.startswith('`['):
            placeholder_text = line.strip('`[] ')
            # Check if next line is a caption
            caption_text = ""
            if i + 1 < n and lines[i+1].strip().startswith('*Caption:'):
                caption_text = lines[i+1].strip('* ')
                i += 1
            add_placeholder_box(doc, placeholder_text, caption_text)
            i += 1
            continue

        # Handle Markdown Tables
        if line.startswith('|') and '|' in line[1:]:
            table_lines = []
            while i < n and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
                
            # Process table
            rows_data = []
            for tline in table_lines:
                # skip separator row
                if re.match(r'^\|[\s\:\-\|]+\|$', tline):
                    continue
                cells = [c.strip() for c in tline.split('|')[1:-1]]
                if cells:
                    rows_data.append(cells)
                    
            if rows_data:
                num_cols = max(len(r) for r in rows_data)
                num_rows = len(rows_data)
                tbl = doc.add_table(rows=num_rows, cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                for r_idx, r_data in enumerate(rows_data):
                    is_header = (r_idx == 0)
                    row = tbl.rows[r_idx]
                    for c_idx, cell_value in enumerate(r_data):
                        if c_idx < len(row.cells):
                            cell = row.cells[c_idx]
                            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            
                            # Clean markdown bold/math from cell text
                            clean_txt = re.sub(r'[\*\$\`]', '', cell_value)
                            run = p.add_run(clean_txt)
                            run.font.name = 'Calibri'
                            run.font.size = Pt(9.5 if not is_header else 10)
                            
                            if is_header:
                                set_cell_background(cell, "0D9488") # Teal header
                                run.bold = True
                                run.font.color.rgb = RGBColor(255, 255, 255)
                            else:
                                if r_idx % 2 == 1:
                                    set_cell_background(cell, "F8FAFC")
                                else:
                                    set_cell_background(cell, "FFFFFF")
                                run.font.color.rgb = RGBColor(15, 23, 42)
                p_after = doc.add_paragraph()
                p_after.paragraph_format.space_after = Pt(6)
            continue

        # Handle Bullet Points
        if line.startswith('* ') or line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            
            raw_text = line[2:].strip()
            # Process inline formatting (bold, math, code)
            parts = re.split(r'(\*\*.*?\*\*|\$.*?\$|\`.*?\`)', raw_text)
            for part in parts:
                if not part:
                    continue
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                elif part.startswith('$') and part.endswith('$'):
                    run = p.add_run(part[1:-1])
                    run.italic = True
                elif part.startswith('`') and part.endswith('`'):
                    run = p.add_run(part[1:-1])
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9.5)
                else:
                    run = p.add_run(part)
                run.font.name = 'Calibri'
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(15, 23, 42)
            i += 1
            continue

        # Handle Numbered Lists
        if re.match(r'^\d+\.\s', line):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            
            raw_text = re.sub(r'^\d+\.\s', '', line).strip()
            parts = re.split(r'(\*\*.*?\*\*|\$.*?\$|\`.*?\`)', raw_text)
            for part in parts:
                if not part:
                    continue
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                elif part.startswith('$') and part.endswith('$'):
                    run = p.add_run(part[1:-1])
                    run.italic = True
                elif part.startswith('`') and part.endswith('`'):
                    run = p.add_run(part[1:-1])
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9.5)
                else:
                    run = p.add_run(part)
                run.font.name = 'Calibri'
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(15, 23, 42)
            i += 1
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        
        parts = re.split(r'(\*\*.*?\*\*|\$.*?\$|\`.*?\`)', line)
        for part in parts:
            if not part:
                continue
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            elif part.startswith('$') and part.endswith('$'):
                run = p.add_run(part[1:-1])
                run.italic = True
            elif part.startswith('`') and part.endswith('`'):
                run = p.add_run(part[1:-1])
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
            else:
                run = p.add_run(part)
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(15, 23, 42)
            
        i += 1

    doc.save(docx_path)
    print(f"Successfully generated Word Document at: {docx_path}")

if __name__ == '__main__':
    md_file = r"c:\Users\Mohsin Ali\Desktop\Data\PEARLS_AQI_Predictor_Final_Project_Report.md"
    word_file = r"c:\Users\Mohsin Ali\Desktop\Data\PEARLS_AQI_Predictor_Final_Project_Report.docx"
    build_word_document(md_file, word_file)
